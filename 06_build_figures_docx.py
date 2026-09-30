# -*- coding: utf-8 -*-
"""Build output/review_figures.docx - figures + data tables for the review pack."""
import csv
import os
import numpy as np
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

here = os.path.dirname(os.path.abspath(__file__))
data_dir = os.path.join(here, '..', 'data')
out_dir = os.path.join(here, '..', 'output')
VARIANT = 'n24_sig'

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(10)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def set_cn(run, cn='宋体', en='Times New Roman'):
    run.font.name = en
    run.element.rPr.rFonts.set(qn('w:eastAsia'), cn)

def para(text, bold=False, size=10, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    set_cn(r)
    if align:
        p.alignment = align
    return p

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htext in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(htext)
        r.bold = True
        r.font.size = Pt(8.5)
        set_cn(r, '黑体')
        tcPr = c._tc.get_or_add_tcPr()
        shd = tcPr.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): 'D9E2F3'})
        tcPr.append(shd)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.rows[ri + 1].cells[ci]
            c.text = ''
            r = c.paragraphs[0].add_run(str(val))
            r.font.size = Pt(8)
            set_cn(r)
    if widths:
        for ci, w in enumerate(widths):
            for row in t.rows:
                row.cells[ci].width = Cm(w)
    return t

# ---------- header ----------
para('Reproducibility Pack - Figures and Data Tables', bold=True, size=16,
     align=WD_ALIGN_PARAGRAPH.CENTER)
para('TEM wavelet-denoising reference pipeline & independent audit of manuscript Table 2',
     size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
para('Companion to Review Comments #1 and #4 - APR Review Report R2, 2026-09-30 | '
     'noise variant: %s (signal-scaled noise level implied by the manuscript\'s '
     'reported all-window noisy SNR of 12.4 dB)' % VARIANT, size=8.5,
     align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()

# ---------- Fig 1 ----------
para('Figure 1. Synthetic TEM decay curve and reference denoising result',
     bold=True, size=11)
para('Denoising configuration per manuscript Table 1: sym8, 5-level DWT, Rigrsure '
     'threshold, improved soft threshold (with classical soft threshold as baseline). '
     'Denoising is applied to the 2048-point time series and then windowed back to '
     'the 40 gates. clean = late-time power-law truth (100 Ω·m half-space).', size=8.5)
doc.add_picture(os.path.join(out_dir, 'fig1_decay.png'), width=Cm(14.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
para('Table 1a. Underlying data of Figure 1 (40 gates, excerpt: first 6 and last 4)',
     bold=True, size=9.5)
t40, clean, noisy, den_i, den_s = [], [], [], [], []
with open(os.path.join(data_dir, f'denoised_{VARIANT}.csv')) as f:
    r = csv.DictReader(f)
    for row in r:
        t40.append(float(row['t_center_ms']))
        clean.append(float(row['clean_uVA']))
        noisy.append(float(row['noisy_uVA']))
        den_i.append(float(row['denoised_improved_uVA']))
        den_s.append(float(row['denoised_soft_uVA']))
rows = []
idx = list(range(6)) + list(range(36, 40))
for i in idx:
    rows.append([i + 1, f'{t40[i]:.4f}', f'{clean[i]:.3e}', f'{noisy[i]:.3e}',
                 f'{den_i[i]:.3e}', f'{den_s[i]:.3e}'])
table(['gate', 't (ms)', 'clean (uV/A)', 'noisy (uV/A)',
       'denoised-improved (uV/A)', 'denoised-soft (uV/A)'],
      rows, widths=[1.2, 1.8, 2.6, 2.6, 3.0, 3.0])
para('Full 40-gate table: data/denoised_%s.csv' % VARIANT, size=8)
doc.add_paragraph()

# ---------- Fig 2 ----------
para('Figure 2. Recomputed quality metrics by window (reference pipeline)', bold=True, size=11)
para('Windows identical to manuscript Table 2. Metric definitions (standard in the '
     'denoising literature): SNR = 10·log10(Σs²/Σ(s-ŝ)²) dB; RMSE = √mean((s-ŝ)²); '
     'MRE = max(|s-ŝ|/|s|)·100 %. Note the modest gain (+1.4 dB all-window SNR) '
     'relative to the manuscript\'s claimed +16.3 dB.', size=8.5)
doc.add_picture(os.path.join(out_dir, 'fig2_metrics.png'), width=Cm(16.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
para('Table 2a. Recomputed metric values (reference pipeline)', bold=True, size=9.5)
WINDOWS = {'early (1-12)': (0, 12), 'middle (13-28)': (12, 28),
           'late (29-40)': (28, 40), 'all (1-40)': (0, 40)}
def snr(s, sh):
    e = np.sum((s - sh) ** 2)
    return 10 * np.log10(np.sum(s ** 2) / e)
def rmse(s, sh):
    return float(np.sqrt(np.mean((s - sh) ** 2)))
def mre(s, sh):
    return float(np.max(np.abs(s - sh) / np.maximum(np.abs(s), 1e-30)) * 100)
clean = np.array(clean); noisy = np.array(noisy)
den_i = np.array(den_i); den_s = np.array(den_s)
rows = []
for lab, (a, b) in WINDOWS.items():
    cs, ns, di, ds = clean[a:b], noisy[a:b], den_i[a:b], den_s[a:b]
    rows.append([lab, 'noisy', f'{snr(cs, ns):.2f}', f'{rmse(cs, ns):.4f}', f'{mre(cs, ns):.2f}'])
    rows.append(['', 'denoised-improved', f'{snr(cs, di):.2f}', f'{rmse(cs, di):.4f}', f'{mre(cs, di):.2f}'])
    rows.append(['', 'denoised-soft', f'{snr(cs, ds):.2f}', f'{rmse(cs, ds):.4f}', f'{mre(cs, ds):.2f}'])
table(['window', 'state', 'SNR (dB)', 'RMSE (uV/A)', 'MRE (%)'], rows,
      widths=[3.0, 3.4, 2.4, 2.6, 2.2])
doc.add_paragraph()

# ---------- Fig 3 ----------
para('Figure 3. Independent arithmetic audit of manuscript Table 2', bold=True, size=11)
para('Panel A: manuscript all-window RMSE vs the value implied by window-weighted '
     'aggregation of its own per-window numbers. Panel B: reported SNR gain vs the '
     'gain implied by the RMSE ratio (these must agree under standard definitions).',
     size=8.5)
doc.add_picture(os.path.join(out_dir, 'fig3_audit.png'), width=Cm(16.0))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
para('Table 3a. Audit summary (manuscript numbers taken as printed)', bold=True, size=9.5)
table(['Check', 'Manuscript', 'Independent recomputation', 'Consistent?'],
      [['All-window RMSE (noisy)', '0.142', '0.180 (window-weighted)', 'No'],
       ['All-window RMSE (denoised)', '0.015', '0.0235 (window-weighted)', 'No'],
       ['All-window SNR gain', '+16.3 dB', '+19.5 dB (implied by RMSE ratio)', 'No (−3.2 dB)'],
       ['Implied noise from noisy SNR 12.4 dB', 'claim: 5 %', '24 % (energy definition)', 'No']],
      widths=[5.2, 3.0, 5.2, 2.4])
para('Full audit output: output/audit_report_manuscript_table2.txt. Statistical '
     'note: with 91 % of signal energy in the first three gates, energy-weighted '
     'SNR over 40 gates is realisation-sensitive (E[SNR]=13.9 dB, std 3.8 dB at the '
     'implied 24 % noise level, Monte-Carlo 2000 draws).', size=8)
doc.add_paragraph()

# ---------- side-by-side ----------
para('Table 4a. Side-by-side: manuscript Table 2 vs reference pipeline (all window)',
     bold=True, size=9.5)
table(['Metric', 'Manuscript noisy', 'Manuscript denoised', 'Ref. noisy', 'Ref. denoised-improved'],
      [['SNR (dB)', '12.4', '28.7', '17.9*', '19.3'],
       ['RMSE (uV/A)', '0.142', '0.015', '0.0148', '0.0126'],
       ['MRE (%)', '68.2', '8.7', '41.5', '15.5']],
      widths=[3.2, 3.0, 3.2, 2.6, 3.4])
para('* Single noise realisation; see the statistical note above. The comparison is '
     'illustrative, not a claim that the manuscript\'s numbers are wrong - it shows '
     'they are not currently checkable without definitions, raw data and scripts.',
     size=8)

doc.save(os.path.join(out_dir, 'review_figures.docx'))
print('SAVED:', os.path.join(out_dir, 'review_figures.docx'),
      os.path.getsize(os.path.join(out_dir, 'review_figures.docx')), 'bytes')
